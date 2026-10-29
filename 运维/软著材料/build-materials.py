# -*- coding: utf-8 -*-
"""
软著申报材料 Word 版批量生成（配套 md2docx.py 的统一版式规范）

产出（全部 .docx，并复制出 .pdf 便于打印/上传）：
    01_申请表填写底稿.docx
    02-03_源程序鉴别材料（60页）.docx      ← 每页 50 行、页眉含名称+版本号+真实页码
    04_软件说明书（…V1.0）.docx            ← 封面 + 页眉 + 页码（提交用）
    05_材料清单与办理流程.docx
    06_身份证明与权利归属模板.docx
    07_演示与运行证明.docx

源程序 Word 版按"每页一节"生成的原因：
    提交的 60 页取自 322 页清单的第 1–30 页与第 293–322 页，Word 的 PAGE 域只能顺序编号 1–60，
    与真实页序不符。因此每一页单独成一个节（section），页眉写入该页的真实页码，
    既满足 50 行/页，又保证页眉页码等于材料真实页序，与已生成的 PDF 完全一致。

用法：python 运维/软著材料/build-materials.py
"""
import os
import shutil
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor

sys.path.insert(0, str(Path(__file__).resolve().parent))
from md2docx import add_field, convert, set_run, shade  # noqa: E402

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "软著申报V1.0"
ASSETS = OUT / "assets"
SOFT = "石油天然气工程咨询成果智能审核系统"
VER = "V1.0"
LINES_PER_PAGE = 50
SRC_ROWS = OUT / "源程序-全文.txt"
SRC_TOTAL_PAGES = 322          # 完整清单页数（见 源程序-统计.json）
HEAD_PAGES, TAIL_PAGES = 30, 30
SOFFICE = next((p for p in [r"C:\Program Files\LibreOffice\program\soffice.exe",
                            r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"] if os.path.exists(p)), None)


def to_pdf(docx: Path):
    if not SOFFICE:
        return None
    subprocess.run([SOFFICE, "--headless", "--norestore", "--convert-to", "pdf",
                    "--outdir", str(docx.parent), str(docx)], check=True, timeout=600,
                   env={**os.environ, "HOME": os.environ.get("TEMP", str(docx.parent))})
    return docx.with_suffix(".pdf")


# ────────────────── ① 普通材料：Markdown → Word ──────────────────
DOCS = [
    ("01_申请表填写底稿.md", "01_申请表填写底稿.docx",
     f"{SOFT} {VER} · 软件著作权登记申请表填写底稿", False),
    ("05_材料清单与办理流程.md", "05_材料清单与办理流程.docx",
     f"{SOFT} {VER} · 材料清单与办理流程", False),
    ("06_身份证明与权利归属模板.md", "06_身份证明与权利归属模板.docx",
     f"{SOFT} {VER} · 身份证明与权利归属（含盖章模板）", False),
    ("07_演示与运行证明.md", "07_演示与运行证明.docx",
     f"{SOFT} {VER} · 演示与运行证明", False),
    ("04_说明书正文.md", f"04_软件说明书（{SOFT}{VER}）.docx",
     f"{SOFT} {VER} · 软件说明书", True),
]


# ────────────────── ② 源程序鉴别材料：固定 50 行/页的 Word ──────────────────
def build_source_docx():
    if not SRC_ROWS.exists():
        print("  ! 缺少 源程序-全文.txt（先跑 build-source.mjs）")
        return None
    rows = SRC_ROWS.read_text(encoding="utf-8").split("\n")
    while rows and rows[-1] == "":
        rows.pop()
    pages = []
    for p in range(1, HEAD_PAGES + 1):
        pages.append((p, rows[(p - 1) * LINES_PER_PAGE: p * LINES_PER_PAGE]))
    tail_start = SRC_TOTAL_PAGES - TAIL_PAGES + 1
    for p in range(tail_start, SRC_TOTAL_PAGES + 1):
        pages.append((p, rows[(p - 1) * LINES_PER_PAGE: p * LINES_PER_PAGE]))

    doc = Document()
    for idx, (pg, lines) in enumerate(pages):
        sec = doc.sections[0] if idx == 0 else doc.add_section(WD_SECTION.NEW_PAGE)
        sec.page_width, sec.page_height = Mm(210), Mm(297)
        sec.top_margin, sec.bottom_margin = Cm(1.6), Cm(1.4)
        sec.left_margin, sec.right_margin = Cm(1.5), Cm(1.2)
        sec.header_distance, sec.footer_distance = Cm(0.8), Cm(0.8)
        # 页眉：软件名称 + 版本号 + 该页真实页码
        sec.header.is_linked_to_previous = False
        hp = sec.header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
        hp.paragraph_format.tab_stops.add_tab_stop(Cm(18.3), 2)  # 2 = RIGHT
        set_run(hp.add_run(f"{SOFT} {VER}"), cn="宋体", size=9)
        hp.add_run("\t")
        set_run(hp.add_run(f"第 {pg} 页 / 共 {SRC_TOTAL_PAGES} 页"), cn="宋体", size=9)
        # 页脚留空（页码已在页眉）
        sec.footer.is_linked_to_previous = False
        if idx > 0:
            sec.header.paragraphs[0]._p.addprevious(OxmlElement("w:p"))  # 占位，避免行距异常
        # 正文：50 行，等宽+宋体，固定行距，无段前段后
        for ln in lines:
            p = doc.add_paragraph()
            pf = p.paragraph_format
            pf.space_before = pf.space_after = Pt(0)
            pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            pf.line_spacing = Pt(13)
            set_run(p.add_run(ln if ln else " "), cn="宋体", en="Consolas", size=8.5)
    dst = OUT / "02-03_源程序鉴别材料（60页）.docx"
    doc.save(str(dst))
    print(f"  → {dst.name}（{len(pages)} 页 × {LINES_PER_PAGE} 行）")
    return dst


if __name__ == "__main__":
    print("生成 Word 版材料（统一版式规范）…")
    for src, dst, header, cover in DOCS:
        p = OUT / src
        if not p.exists():
            print(f"  ! 缺 {src}，跳过")
            continue
        out = convert(p, OUT / dst, header=header, cover=cover, cover_title=SOFT, assets=ASSETS)
        print(f"  → {out.name}")
    build_source_docx()

    if "--no-pdf" not in sys.argv:
        print("转 PDF（打印/上传备用）…")
        for dst in [d for _, d, _, _ in DOCS] + ["02-03_源程序鉴别材料（60页）.docx"]:
            f = OUT / dst
            if f.exists():
                pdf = to_pdf(f)
                if pdf:
                    print(f"  → {pdf.name}")
    print("完成。")
